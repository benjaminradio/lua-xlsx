"""Generate tests/fixtures/sample.xlsx (requires openpyxl)."""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

wb = Workbook()

# A tidy-ish sheet with a title banner above the real header row.
ws = wb.active
ws.title = "People"
ws["A1"] = "Quarterly Report"
ws.merge_cells("A1:D1")
for c, h in enumerate(["Name", "Age", "City", "Joined"], 1):
    ws.cell(3, c, h)
rows = [
    ("Alice", 30, "Paris", dt.date(2020, 1, 15)),
    ("Bob", 25.5, "London", dt.datetime(2021, 6, 30, 13, 30)),
    ("Carol", None, "Berlin", dt.date(2022, 12, 1)),
]
for r, row in enumerate(rows, 4):
    for c, v in enumerate(row, 1):
        ws.cell(r, c, v)
ws["F4"] = True
ws["F5"] = False
ws["G4"] = "=B4+1"
ws["H4"] = "#DIV/0!"
ws["A8"] = "Total"
ws["B8"] = "=SUM(B4:B6)"

# Semi-structured sheet: vertically merged group labels.
m = wb.create_sheet("Messy")
m["A1"] = "Region"
m["B1"] = "Item"
m["C1"] = "Qty"
m["A2"] = "North"
m.merge_cells("A2:A4")
for i, (item, q) in enumerate([("nut", 1), ("bolt", 2), ("screw", 3)], 2):
    m.cell(i, 2, item)
    m.cell(i, 3, q)
m["A5"] = "South"
m.merge_cells("A5:A6")
m["B5"] = "washer"
m["C5"] = 10
m["B6"] = "nail"
m["C6"] = 20
m["E2"] = "Notes:"
m["F2"] = "see  appendix"
m["B10"] = "Name"  # duplicated header names below
m["C10"] = "Name"
m["D10"] = None
m["E10"] = "Name"
m["B11"] = 1
m["C11"] = 2
m["D11"] = 3
m["E11"] = 4

# An Excel table.
t = wb.create_sheet("Tbl")
t.append(["Region", "Units", "Price"])
t.append(["EU", 5, 1.5])
t.append(["US", 7, 2.25])
t.append(["APAC", 9, 3.0])
t.add_table(Table(displayName="Sales", ref="A1:C4"))

h = wb.create_sheet("Hidden")
h.sheet_state = "hidden"
h["A1"] = "secret"

wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="People!$B$8")

os.makedirs("tests/fixtures", exist_ok=True)
wb.save("tests/fixtures/sample.xlsx")
